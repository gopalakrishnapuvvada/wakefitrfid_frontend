"""Read the client FG master data template without changing the workbook."""

from __future__ import annotations

from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook

from schemas.master_data import is_valid_image_url


TEMPLATE_HEADERS = (
    "S.no", "Material Code", "Part Number", "Category", "Model",
    "Product Description", "Dimensions (LxWxH mm)", "Color", "Status",
    "FG Image 1", "FG Image 2", "FG Image 3", "FG Image 4",
)
MAX_IMPORT_ROWS = 5000
DIMENSIONS = re.compile(r"^\s*(\d+)\s*[x×]\s*(\d+)\s*[x×]\s*(\d+)\s*(?:mm)?\s*$", re.IGNORECASE)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value: Any) -> str:
    return " ".join(_text(value).casefold().split())


def parse_master_data_workbook(contents: bytes) -> tuple[list[dict], list[dict], int]:
    """Return valid rows, row errors, and the count of nonempty data rows."""
    workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as exc:
            raise ValueError("The workbook is empty.") from exc

        positions = {_header(value): index for index, value in enumerate(headers) if _header(value)}
        missing = [name for name in TEMPLATE_HEADERS if _header(name) not in positions]
        if missing:
            raise ValueError("Missing template columns: " + ", ".join(missing))

        def cell(values: tuple, name: str) -> str:
            index = positions[_header(name)]
            return _text(values[index]) if index < len(values) else ""

        valid_rows: list[dict] = []
        issues: list[dict] = []
        total_rows = 0
        for row_number, values in enumerate(rows, start=2):
            if not any(_text(value) for value in values):
                continue
            total_rows += 1
            if total_rows > MAX_IMPORT_ROWS:
                raise ValueError(f"A workbook can contain at most {MAX_IMPORT_ROWS} data rows.")

            errors: list[str] = []
            fields = {name: cell(values, name) for name in TEMPLATE_HEADERS}
            for name in ("Material Code", "Part Number", "Category", "Model", "Color", "Status"):
                if not fields[name]:
                    errors.append(f"{name} is required")
            for name in ("Material Code", "Part Number", "Category", "Model", "Color"):
                if len(fields[name]) > 100:
                    errors.append(f"{name} exceeds 100 characters")

            dimensions = DIMENSIONS.fullmatch(fields["Dimensions (LxWxH mm)"])
            if not dimensions:
                errors.append("Dimensions must be LxWxH in millimetres, for example 1981x1828x203")

            status_key = re.sub(r"[\s_-]+", "", fields["Status"].casefold())
            status_name = {"active": "Active", "onhold": "On hold", "inactive": "Inactive"}.get(status_key)
            if fields["Status"] and status_name is None:
                errors.append("Status must be Active, OnHold/On hold, or Inactive")

            images: list[str] = []
            for index in range(1, 5):
                url = fields[f"FG Image {index}"]
                if url and not is_valid_image_url(url):
                    errors.append(f"FG Image {index} must be an HTTP/HTTPS URL or a /relative/path")
                elif url and url not in images:
                    images.append(url)

            if errors:
                issues.append({"row": row_number, "kind": "error", "message": "; ".join(errors)})
                continue

            length, width, height = (int(value) for value in dimensions.groups())
            valid_rows.append({
                "row": row_number,
                "material_code": fields["Material Code"],
                "part_number": fields["Part Number"],
                "category": fields["Category"],
                "model": fields["Model"],
                "product_description": fields["Product Description"] or None,
                "length_mm": length,
                "width_mm": width,
                "height_mm": height,
                "color": fields["Color"],
                "status": status_name,
                "fg_image": images,
            })

        if total_rows == 0:
            raise ValueError("The workbook contains no FG master data rows.")
        return valid_rows, issues, total_rows
    finally:
        workbook.close()
